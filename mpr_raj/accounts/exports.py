"""
Exports for one reporting month (plan §5).

`tables(period)` gathers the month once — every section plus the project figures —
as (title, headers, rows of plain strings). The four writers below only format it;
none of them touches the ORM. A new format is one function and one WRITERS entry.
"""
import csv
import datetime
import io
import os
import re
import sys
from pathlib import Path

from django.contrib.staticfiles import finders
from django.template.loader import render_to_string
from django.utils import timezone

from .models import MPREntry, ParameterValue


def _text(value):
    """Every cell reaches the writers as a string — the figures are text already."""
    if value is None or value == "":
        return ""
    if isinstance(value, datetime.date):
        return value.strftime("%d-%m-%Y")
    return str(value)


PARAMETERS = "parameters"


def tables(period, only=None):
    """
    The month as a list of {key, title, headers, rows} — one block per category, plus
    the project figures. Gathered once and used twice: the on-screen reports render a
    block as HTML, the writers below turn the same block into a file. `only` narrows it
    to a single category so a report and its download can't drift apart.
    """
    entries = list(period.entries.select_related("author", "project"))
    out = []
    for kind, label in MPREntry.KIND_CHOICES:
        fields = MPREntry.FIELDS[kind]
        out.append({
            "key": kind,
            "title": label,
            # Project isn't a form field any more — it's the report the row was filed
            # in — so it's a column here rather than one of FIELDS.
            "headers": ["Filed by", "Project"] + [str(MPREntry.label(kind, f)) for f in fields],
            "rows": [[_text(e.author.name or e.author.username), _text(e.project)]
                     + [_text(e.value(f)) for f in fields]
                     for e in entries if e.kind == kind],
        })

    values = (ParameterValue.objects
              .filter(period=period)
              .select_related("parameter", "parameter__project", "parameter__project__leader")
              .order_by("parameter__project__name", "parameter__order", "parameter__name"))
    rows = []
    for value in values:
        project = value.parameter.project
        leader = project.leader
        rows.append([project.code, project.name, project.prism_id or "",
                     leader.name or leader.username if leader else "",
                     value.parameter.name,
                     value.previous_month, value.reporting_month, value.cumulative])
    out.append({
        "key": PARAMETERS,
        "title": "Project parameters",
        "headers": ["Project ID", "Project", "PRISM ID", "Project Leader", "Parameter",
                    "Previous month", "Reporting month", "Cumulative since inception"],
        "rows": rows,
    })
    return [t for t in out if t["key"] == only] if only else out


def _csv(period, tabs):
    # One file, one block per section — a workbook's worth of sheets doesn't fit CSV.
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([f"Monthly Project Report — {period}"])
    for table in tabs:
        writer.writerow([])
        writer.writerow([table["title"]])
        writer.writerow(table["headers"])
        writer.writerows(table["rows"] or [["No entries."]])
    # utf-8-sig: Excel on Windows otherwise mangles ₹ and Devanagari.
    return buf.getvalue().encode("utf-8-sig"), "text/csv"


def _xlsx(period, tabs):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    wb.remove(wb.active)
    for table in tabs:
        # Excel sheet names: 31 chars, and none of []:*?/\
        ws = wb.create_sheet(re.sub(r"[\[\]:*?/\\]", "-", table["title"])[:31])
        ws.append(table["headers"])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in table["rows"]:
            ws.append(row)
        ws.freeze_panes = "A2"
        for i, header in enumerate(table["headers"], start=1):
            longest = max([len(header)] + [len(r[i - 1]) for r in table["rows"]])
            ws.column_dimensions[ws.cell(1, i).column_letter].width = min(max(longest + 2, 12), 50)
            ws.cell(1, i).alignment = Alignment(vertical="top", wrap_text=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _docx(period, tabs):
    from docx import Document

    doc = Document()
    doc.add_heading(f"Monthly Project Report — {period}", 0)
    doc.add_paragraph("National Informatics Centre, Rajasthan")
    for table in tabs:
        doc.add_heading(table["title"], level=1)
        if not table["rows"]:
            doc.add_paragraph("No entries.")
            continue
        grid = doc.add_table(rows=1, cols=len(table["headers"]))
        grid.style = "Table Grid"
        for cell, header in zip(grid.rows[0].cells, table["headers"], strict=True):
            cell.text = header
            cell.paragraphs[0].runs[0].bold = True
        for row in table["rows"]:
            for cell, value in zip(grid.add_row().cells, row, strict=True):
                cell.text = value
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _dyld_fallback():
    """
    macOS: WeasyPrint dlopen()s Pango/Cairo by leaf name and dyld doesn't search
    Homebrew's prefix, so the import fails unless the server happened to be started
    with DYLD_FALLBACK_LIBRARY_PATH set. Setting it here, before the import, works and
    doesn't depend on how anyone launched runserver. No-op everywhere else.
    """
    if sys.platform != "darwin":
        return
    # An unset variable means dyld's built-in chain, which we must keep.
    default = f"{os.path.expanduser('~')}/lib:/usr/local/lib:/lib:/usr/lib"
    paths = (os.environ.get("DYLD_FALLBACK_LIBRARY_PATH") or default).split(":")
    for brew in ("/opt/homebrew/lib", "/usr/local/lib"):
        if brew not in paths:
            paths.append(brew)
    os.environ["DYLD_FALLBACK_LIBRARY_PATH"] = ":".join(paths)


def _pdf(period, tabs):
    _dyld_fallback()
    from weasyprint import HTML

    logo = finders.find("accounts/nic-logo.png")
    fonts = finders.find("accounts/fonts/inter-latin.woff2")
    html = render_to_string("accounts/export.html", {
        "period": period,
        "tables": tabs,
        "logo": Path(logo).as_uri() if logo else None,
        "fonts": Path(fonts).parent.as_uri() if fonts else "",
        "generated": timezone.localdate(),
    })
    return HTML(string=html).write_pdf(), "application/pdf"


WRITERS = {"csv": _csv, "xlsx": _xlsx, "docx": _docx, "pdf": _pdf}


def render(period, fmt, only=None):
    """(bytes, content_type) for one month in one format, or one category of it."""
    return WRITERS[fmt](period, tables(period, only))
